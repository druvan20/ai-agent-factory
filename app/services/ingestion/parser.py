from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

import fitz  # type: ignore[import-untyped]  # PyMuPDF
from docx import Document as DocxDocument  # type: ignore[import-untyped]
from pptx import Presentation  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]


SUPPORTED_MIMES: dict[str, str] = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/markdown": "md",
    "text/plain": "txt",
}

EXTENSION_TO_MIME: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".md": "text/markdown",
    ".txt": "text/plain",
}


@dataclass
class ParsedSection:
    heading: str
    level: int  # 0 = root / no heading
    content: str
    page_or_slide: int | None = None


@dataclass
class ParsedDocument:
    filename: str
    mime_type: str
    sections: list[ParsedSection] = field(default_factory=list)
    raw_text: str = ""


def parse_document(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    kind = SUPPORTED_MIMES.get(mime_type)
    if kind is None:
        raise ValueError(f"Unsupported MIME type: {mime_type}")

    dispatch: dict[str, Any] = {
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
        "xlsx": _parse_xlsx,
        "md": _parse_markdown,
        "txt": _parse_txt,
    }
    return dispatch[kind](content, filename, mime_type)


def _parse_pdf(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    doc: Any = fitz.open(stream=content, filetype="pdf")
    sections: list[ParsedSection] = []
    full_text_parts: list[str] = []

    for page_num in range(len(doc)):
        page: Any = doc[page_num]
        text_dict: Any = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
        blocks: list[Any] = text_dict["blocks"]
        page_sections = _extract_heading_sections_from_blocks(blocks, page_num + 1)
        sections.extend(page_sections)
        full_text_parts.append(str(page.get_text()))

    doc.close()

    if not sections:
        sections.append(ParsedSection(heading="Document", level=0, content="\n".join(full_text_parts)))

    return ParsedDocument(
        filename=filename,
        mime_type=mime_type,
        sections=sections,
        raw_text="\n".join(full_text_parts),
    )


def _extract_heading_sections_from_blocks(
    blocks: list[Any], page_num: int
) -> list[ParsedSection]:
    sections: list[ParsedSection] = []
    current_heading = "Untitled"
    current_level = 0
    current_lines: list[str] = []

    for block in blocks:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans: list[Any] = line.get("spans", [])
            text = "".join(str(span["text"]) for span in spans).strip()
            if not text:
                continue
            avg_size: float = (
                sum(float(s["size"]) for s in spans) / len(spans)
                if spans else 12.0
            )
            is_bold = all("bold" in str(s.get("font", "")).lower() for s in spans)

            if avg_size >= 14 or (avg_size >= 12 and is_bold):
                if current_lines:
                    sections.append(
                        ParsedSection(
                            heading=current_heading,
                            level=current_level,
                            content="\n".join(current_lines),
                            page_or_slide=page_num,
                        )
                    )
                level = 1 if avg_size >= 18 else (2 if avg_size >= 14 else 3)
                current_heading = text
                current_level = level
                current_lines = []
            else:
                current_lines.append(text)

    if current_lines:
        sections.append(
            ParsedSection(heading=current_heading, level=current_level, content="\n".join(current_lines), page_or_slide=page_num)
        )

    return sections


def _parse_docx(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    doc: Any = DocxDocument(io.BytesIO(content))
    sections: list[ParsedSection] = []
    current_heading = "Document"
    current_level = 0
    current_lines: list[str] = []

    for para in doc.paragraphs:
        text: str = para.text.strip()
        if not text:
            continue

        style: Any = para.style
        style_name: str = str(style.name or "").lower() if style else ""
        if style_name.startswith("heading"):
            if current_lines:
                sections.append(ParsedSection(heading=current_heading, level=current_level, content="\n".join(current_lines)))
            try:
                current_level = int(style_name.replace("heading", "").strip())
            except ValueError:
                current_level = 1
            current_heading = text
            current_lines = []
        else:
            current_lines.append(text)

    if current_lines:
        sections.append(ParsedSection(heading=current_heading, level=current_level, content="\n".join(current_lines)))

    raw_text = "\n".join(str(p.text) for p in doc.paragraphs)
    return ParsedDocument(filename=filename, mime_type=mime_type, sections=sections or [ParsedSection(heading="Document", level=0, content=raw_text)], raw_text=raw_text)


def _parse_pptx(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    prs: Any = Presentation(io.BytesIO(content))
    sections: list[ParsedSection] = []
    full_parts: list[str] = []

    for slide_num, slide in enumerate(prs.slides, 1):
        title = ""
        body_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                tf: Any = shape.text_frame
                text: str = str(tf.text).strip()
                if shape == slide.shapes.title and text:
                    title = text
                elif text:
                    body_parts.append(text)
        body = "\n".join(body_parts)
        full_parts.append(f"{title}\n{body}" if title else body)
        if title or body:
            sections.append(ParsedSection(heading=title or f"Slide {slide_num}", level=1, content=body, page_or_slide=slide_num))

    return ParsedDocument(filename=filename, mime_type=mime_type, sections=sections, raw_text="\n\n".join(full_parts))


def _parse_xlsx(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    wb: Any = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sections: list[ParsedSection] = []
    full_parts: list[str] = []

    for sheet_name in wb.sheetnames:
        ws: Any = wb[sheet_name]
        rows_text: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            line = " | ".join(cells)
            if line.strip(" |"):
                rows_text.append(line)
        body = "\n".join(rows_text)
        full_parts.append(body)
        if body:
            sections.append(ParsedSection(heading=str(sheet_name), level=1, content=body))

    wb.close()
    return ParsedDocument(filename=filename, mime_type=mime_type, sections=sections, raw_text="\n\n".join(full_parts))


def _parse_markdown(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace")
    sections: list[ParsedSection] = []
    current_heading = "Document"
    current_level = 0
    current_lines: list[str] = []

    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.*)", line)
        if match:
            if current_lines:
                sections.append(ParsedSection(heading=current_heading, level=current_level, content="\n".join(current_lines)))
            current_level = len(match.group(1))
            current_heading = match.group(2).strip()
            current_lines = []
        else:
            current_lines.append(line)

    if current_lines:
        sections.append(ParsedSection(heading=current_heading, level=current_level, content="\n".join(current_lines)))

    return ParsedDocument(filename=filename, mime_type=mime_type, sections=sections, raw_text=text)


def _parse_txt(content: bytes, filename: str, mime_type: str) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace")
    return ParsedDocument(
        filename=filename,
        mime_type=mime_type,
        sections=[ParsedSection(heading="Document", level=0, content=text)],
        raw_text=text,
    )